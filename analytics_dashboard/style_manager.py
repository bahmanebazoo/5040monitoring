"""
مدیریت استایل اکسل — SRP
"""

from __future__ import annotations
from typing import Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class StyleManager:

    HEADER_FILL = PatternFill(
        start_color="1F4E79", end_color="1F4E79", fill_type="solid"
    )
    HEADER_FONT = Font(name="B Nazanin", size=12, bold=True, color="FFFFFF")
    KPI_FILL = PatternFill(
        start_color="D6EAF8", end_color="D6EAF8", fill_type="solid"
    )
    GOOD_FILL = PatternFill(
        start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
    )
    BAD_FILL = PatternFill(
        start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
    )
    NORMAL_FONT = Font(name="B Nazanin", size=11)
    SECTION_FONT = Font(name="B Nazanin", size=13, bold=True, color="1F4E79")
    THIN_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # ──────────────────────────────────────
    #  هدر
    # ──────────────────────────────────────

    @classmethod
    def apply_header(
        cls, ws: Worksheet, row: int = 1, max_col: Optional[int] = None
    ):
        if max_col is None:
            max_col = ws.max_column
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = cls.THIN_BORDER

    # ──────────────────────────────────────
    #  داده (با محدوده مشخص)
    # ──────────────────────────────────────

    @classmethod
    def style_data(
        cls,
        ws: Worksheet,
        start_row: int = 2,
        end_row: Optional[int] = None,
        max_col: Optional[int] = None,
    ):
        """
        Border + Font + Alignment برای ردیف‌های داده.
        اگر end_row/max_col داده نشود، تا آخر شیت اعمال می‌شود.
        """
        if end_row is None:
            end_row = ws.max_row
        if max_col is None:
            max_col = ws.max_column

        for row in ws.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=1,
            max_col=max_col,
        ):
            for cell in row:
                cell.font = cls.NORMAL_FONT
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                cell.border = cls.THIN_BORDER

    # ──────────────────────────────────────
    #  عنوان بخش (Merge & Center)
    # ──────────────────────────────────────

    @classmethod
    def write_section_title(
        cls,
        ws: Worksheet,
        row: int,
        title: str,
        col_count: int,
    ) -> int:
        """
        عنوان بخش را در یک ردیف merge‌شده می‌نویسد.

        بازگشت: شماره ردیف بعد از عنوان (row + 1)
        """
        end_col_letter = get_column_letter(col_count)
        ws.merge_cells(f"A{row}:{end_col_letter}{row}")
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = cls.SECTION_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cls.THIN_BORDER
        return row + 1

    # ──────────────────────────────────────
    #  عرض خودکار ستون‌ها
    # ──────────────────────────────────────

    @classmethod
    def auto_fit(cls, ws: Worksheet):
        for col_cells in ws.columns:
            max_len = 0
            letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(max_len + 4, 45)
