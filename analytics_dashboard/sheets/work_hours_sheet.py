"""
شیت تحلیل ساعت کاری — داخل تایم vs خارج تایم
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.merge import MergedCell

from ..data_preparator import PreparedData
from .base import SheetCreator


class WorkHoursSheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "تحلیل_ساعت_کاری"

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df

        if "_in_work_hours" not in df.columns:
            ws.cell(row=1, column=1, value="داده ساعت کاری موجود نیست")
            return

        in_wh = df[df["_in_work_hours"] == True]
        out_wh = df[df["_in_work_hours"] == False]

        # ═══════════════════════════════════════
        # ① جدول خلاصه کلی
        # ═══════════════════════════════════════
        summary = pd.DataFrame([
            self._calc_row("داخل تایم کاری", in_wh, df, data),
            self._calc_row("خارج تایم کاری", out_wh, df, data),
            self._calc_row("کل", df, df, data),
        ])

        next_row = self._write_table(ws, summary, start_row=1)

        # ═══════════════════════════════════════
        # ② Pie Chart — نسبت داخل/خارج
        # ═══════════════════════════════════════
        pie = PieChart()
        pie.title = "نسبت تماس‌ها — داخل vs خارج تایم کاری"
        pie.style = 10
        pie.width, pie.height = 18, 14
        pie.add_data(
            Reference(ws, min_col=2, min_row=1, max_row=3),
            titles_from_data=True,
        )
        pie.set_categories(
            Reference(ws, min_col=1, min_row=2, max_row=3)
        )
        # ✅ نمایش درصد روی قطعه‌های پای
        if pie.series:
            pie.series[0].dLbls = DataLabelList()
            pie.series[0].dLbls.showPercent = True
            pie.series[0].dLbls.showVal = True
            pie.series[0].dLbls.numFmt = "#,##0"
        ws.add_chart(pie, f"{get_column_letter(len(summary.columns) + 2)}1")

        # ═══════════════════════════════════════
        # ③ جدول تفکیک روزانه
        # ═══════════════════════════════════════
        if "_weekday_num" in df.columns:
            next_row = self._add_daily_breakdown(
                ws, df, data, next_row + 1
            )

        # ═══════════════════════════════════════
        # ④ جدول تفکیک ساعتی (فقط خارج تایم)
        # ═══════════════════════════════════════
        if "_hour" in df.columns and len(out_wh) > 0:
            self._add_out_of_hours_detail(ws, out_wh, next_row + 1)

        self.style.auto_fit(ws)

    # ── ردیف خلاصه ──

    def _calc_row(
        self, label: str, subset: pd.DataFrame,
        full_df: pd.DataFrame, data: PreparedData,
    ) -> dict:
        total = len(subset)
        conn = int(self._connected_mask(subset).sum())
        s_match = int(
            self._match_mask(subset, data.support_match_col).sum()
        )

        row = {
            "بازه": label,
            "تعداد": total,
            "درصد_از_کل": f"{total / max(len(full_df), 1) * 100:.1f}%",
            "وصل_شده": conn,
            "نرخ_پاسخ_%": f"{conn / max(total, 1) * 100:.1f}%",
            "مچ_پشتیبانی": s_match,
        }

        if "_wait_seconds" in subset.columns and total > 0:
            row["میانگین_انتظار"] = round(
                subset["_wait_seconds"].mean(), 1
            )
            row["≥2min"] = int(
                (subset["_wait_seconds"] >= 120).sum()
            )

        if "_talk_seconds" in subset.columns and total > 0:
            row["میانگین_مکالمه"] = round(
                subset["_talk_seconds"].mean(), 1
            )

        return row

    # ── هدر مرج‌شده برای بخش‌ها ──

    def _write_section_header(self, ws, row: int, text: str, num_cols: int):
        """
        ✅ هدر بخش را می‌نویسد و سلول‌ها را به اندازه تعداد ستون جدول مرج می‌کند.
        """
        ws.cell(row=row, column=1, value=text)
        if num_cols > 1:
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row, end_column=num_cols,
            )
        self.style.apply_header(ws, row=row, max_col=num_cols)

    # ── تفکیک روزانه ──

    def _add_daily_breakdown(
        self, ws, df: pd.DataFrame, data: PreparedData, start_row: int,
    ) -> int:
        # مپ weekday_num به نام فارسی روز
        day_names = {
            5: "شنبه",
            6: "یکشنبه",
            0: "دوشنبه",
            1: "سه‌شنبه",
            2: "چهارشنبه",
            3: "پنجشنبه",
            4: "جمعه",
        }

        # ساعت کاری هر روز (برای نمایش)
        work_schedule = {
            5: "8–18", 6: "8–18", 0: "8–18", 1: "8–18", 2: "8–18",
            3: "8–14",
            4: "تعطیل",
        }

        # ترتیب ایرانی: شنبه تا جمعه
        ir_order = [5, 6, 0, 1, 2, 3, 4]

        rows = []
        for wd in ir_order:
            day_df = df[df["_weekday_num"] == wd]
            if day_df.empty:
                continue

            in_wh = day_df[day_df["_in_work_hours"] == True]
            out_wh = day_df[day_df["_in_work_hours"] == False]

            rows.append({
                "روز": day_names[wd],
                "ساعت_کاری": work_schedule[wd],
                "کل": len(day_df),
                "داخل_تایم": len(in_wh),
                "خارج_تایم": len(out_wh),
                "درصد_خارج": f"{len(out_wh) / max(len(day_df), 1) * 100:.1f}%",
            })

        if not rows:
            return start_row

        daily_df = pd.DataFrame(rows)
        ncols = len(daily_df.columns)

        # ✅ عنوان بخش — مرج‌شده به اندازه تعداد ستون‌های جدول
        self._write_section_header(
            ws, start_row,
            "📅 تفکیک روزانه — داخل/خارج تایم کاری",
            ncols,
        )

        next_row = self._write_table(ws, daily_df, start_row=start_row + 1)

        # ── Bar Chart روزانه ──
        n_days = len(rows)
        bar = BarChart()
        bar.type = "col"
        bar.title = "تعداد تماس روزانه — داخل vs خارج تایم"
        bar.style = 10
        bar.width, bar.height = 28, 14

        # ✅ محور Y: مقادیر عددی مشخص
        bar.y_axis.title = "تعداد تماس"
        bar.y_axis.delete = False
        bar.y_axis.numFmt = "#,##0"
        bar.y_axis.tickLblPos = "low"

        # ✅ محور X
        bar.x_axis.title = "روز هفته"
        bar.x_axis.delete = False
        bar.x_axis.tickLblPos = "low"

        bar.add_data(
            Reference(
                ws,
                min_col=4, max_col=5,
                min_row=start_row + 1,
                max_row=start_row + 1 + n_days,
            ),
            titles_from_data=True,
        )
        bar.set_categories(
            Reference(
                ws,
                min_col=1,
                min_row=start_row + 2,
                max_row=start_row + 1 + n_days,
            )
        )

        # ✅ نمایش مقدار روی هر ستون
        for s in bar.series:
            s.dLbls = DataLabelList()
            s.dLbls.showVal = True
            s.dLbls.numFmt = "#,##0"

        ws.add_chart(
            bar,
            f"{get_column_letter(ncols + 2)}{29}"
        )

        return next_row

    # ── جزئیات خارج تایم ──

    def _add_out_of_hours_detail(
        self, ws, out_wh: pd.DataFrame, start_row: int
    ):
        hourly = (
            out_wh.groupby("_hour")
            .size()
            .reset_index(name="تعداد")
        )
        hourly.rename(columns={"_hour": "ساعت"}, inplace=True)
        hourly.sort_values("ساعت", inplace=True)

        ncols = len(hourly.columns)

        # ✅ عنوان بخش — مرج‌شده به اندازه تعداد ستون‌های جدول
        self._write_section_header(
            ws, start_row,
            "🌙 جزئیات تماس‌های خارج تایم کاری (به تفکیک ساعت)",
            ncols,
        )

        self._write_table(ws, hourly, start_row=start_row + 1)
