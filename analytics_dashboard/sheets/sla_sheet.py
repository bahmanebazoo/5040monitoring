"""
شیت SLA
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from ..data_preparator import PreparedData
from .base import SheetCreator


class SLASheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "SLA"

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df

        if (
            "_wait_seconds" not in df.columns
            or "_status" not in df.columns
        ):
            ws.cell(row=1, column=1, value="داده SLA موجود نیست")
            return

        conn_df = df[self._connected_mask(df)].copy()
        total_conn = len(conn_df)

        thresholds = [10, 15, 20, 30, 45, 60, 90, 120]

        # ── هدر جدول اصلی ──
        ws.cell(row=1, column=1, value="آستانه (ثانیه)")
        ws.cell(row=1, column=2, value="تعداد_پاسخ")
        ws.cell(row=1, column=3, value="SLA_%")
        ws.cell(row=1, column=4, value="≥2min")
        self.style.apply_header(ws, max_col=4)

        for ri, th in enumerate(thresholds, 2):
            cnt = int((conn_df["_wait_seconds"] <= th).sum())
            pct = cnt / max(total_conn, 1) * 100

            over_2min = int((conn_df["_wait_seconds"] >= 120).sum())

            ws.cell(row=ri, column=1, value=th)
            ws.cell(row=ri, column=2, value=cnt)
            ws.cell(row=ri, column=3, value=f"{pct:.1f}%")
            ws.cell(row=ri, column=4, value=over_2min)

        # ✅ border جدول اصلی
        self.style.style_data(
            ws, start_row=2, end_row=len(thresholds) + 1, max_col=4
        )

        n = len(thresholds)

        # ── خلاصه آمار ≥2min ──
        summary_row = n + 3
        over_2min_total = int(
            (conn_df["_wait_seconds"] >= 120).sum()
        )
        over_2min_pct = (
            over_2min_total / max(total_conn, 1) * 100
        )

        # ── Line Chart ──
        line = LineChart()
        line.title = "منحنی SLA"
        line.style = 10
        line.width, line.height = 25, 14

        # ✅ محور Y — نمایش اعداد
        line.y_axis.title = "تعداد"
        line.y_axis.delete = False
        line.y_axis.numFmt = "#,##0"
        line.y_axis.tickLblPos = "low"

        # ✅ محور X
        line.x_axis.title = "ثانیه"
        line.x_axis.delete = False
        line.x_axis.tickLblPos = "low"

        line.add_data(
            Reference(ws, min_col=2, min_row=1, max_row=n + 1),
            titles_from_data=True,
        )
        line.set_categories(
            Reference(ws, min_col=1, min_row=2, max_row=n + 1)
        )
        if line.series:
            line.series[0].graphicalProperties.line.width = 25000
        ws.add_chart(line, "H1")

        # ── SLA ساعتی ──
        if "_hour" in conn_df.columns:
            self._add_hourly_sla(ws, conn_df, summary_row + 2)

        self.style.auto_fit(ws)

    # ── جدول ساعتی ──

    def _add_hourly_sla(self, ws, conn_df, sr: int):
        ws.cell(row=sr, column=1, value="ساعت")
        ws.cell(row=sr, column=2, value="کل_وصل")
        ws.cell(row=sr, column=3, value="≤20s")
        ws.cell(row=sr, column=4, value="SLA_%")
        ws.cell(row=sr, column=5, value="≥2min")
        ws.cell(row=sr, column=6, value="≥2min_%")
        self.style.apply_header(ws, row=sr, max_col=6)

        # ✅ فقط ساعت‌هایی که آمار دارند (بدون نمایش ساعت‌های خالی)
        hours_with_data = sorted(conn_df["_hour"].dropna().unique())

        row_idx = sr + 1  # ردیف شروع داده‌ها
        for h in hours_with_data:
            hdf = conn_df[conn_df["_hour"] == h]
            hc = len(hdf)

            # ✅ اگر هیچ تماسی در این ساعت نبود، رد شو
            if hc == 0:
                continue

            s20 = int((hdf["_wait_seconds"] <= 20).sum())
            over2 = int((hdf["_wait_seconds"] >= 120).sum())
            over2_pct = over2 / max(hc, 1) * 100

            ws.cell(row=row_idx, column=1, value=int(h))
            ws.cell(row=row_idx, column=2, value=hc)
            ws.cell(row=row_idx, column=3, value=s20)
            ws.cell(
                row=row_idx, column=4,
                value=f"{s20 / max(hc, 1) * 100:.1f}%",
            )
            ws.cell(row=row_idx, column=5, value=over2)
            ws.cell(row=row_idx, column=6, value=f"{over2_pct:.1f}%")
            row_idx += 1

        data_end_row = row_idx - 1
        num_data_rows = data_end_row - sr  # تعداد ردیف‌های داده

        # ✅ border جدول ساعتی
        if num_data_rows > 0:
            self.style.style_data(
                ws,
                start_row=sr + 1,
                end_row=data_end_row,
                max_col=6,
            )

        # ── Bar Chart ساعتی (شامل ≥2min) ──
        if num_data_rows > 0:
            bar = BarChart()
            bar.type = "col"
            bar.title = "SLA ساعتی (≤20s و ≥2min)"
            bar.style = 10
            bar.width, bar.height = 30, 15

            # ✅ محور Y — نمایش اعداد
            bar.y_axis.delete = False
            bar.y_axis.numFmt = "#,##0"
            bar.y_axis.tickLblPos = "low"
            bar.y_axis.title = "تعداد"

            # ✅ محور X
            bar.x_axis.delete = False
            bar.x_axis.tickLblPos = "low"

            bar.add_data(
                Reference(
                    ws,
                    min_col=2,
                    max_col=5,
                    min_row=sr,
                    max_row=data_end_row,
                ),
                titles_from_data=True,
            )
            bar.set_categories(
                Reference(
                    ws, min_col=1, min_row=sr + 1, max_row=data_end_row
                )
            )
            ws.add_chart(bar, "H26")
