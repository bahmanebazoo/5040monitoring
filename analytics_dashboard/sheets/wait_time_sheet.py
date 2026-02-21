"""
شیت تحلیل زمان انتظار و مکالمه
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.worksheet import Worksheet

from ..data_preparator import PreparedData
from .base import SheetCreator


class WaitTimeSheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "تحلیل_زمان_انتظار"

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df
        has_w = "_wait_seconds" in df.columns
        has_t = "_talk_seconds" in df.columns

        if not has_w and not has_t:
            ws.cell(row=1, column=1, value="داده زمان موجود نیست")
            return

        rc = 1

        if has_w:
            w = df["_wait_seconds"].dropna()
            rc = self._write_stats(ws, w, "آمار انتظار (ثانیه)", rc)
            rc += 1
            rc, ds, dist_len = self._write_dist(
                ws, w, rc,
                bins=[0, 10, 20, 30, 60, 120, 300, float("inf")],
                labels=[
                    "0-10s", "10-20s", "20-30s", "30-60s",
                    "1-2m", "2-5m", "5m+",
                ],
            )
            self._add_bar(ws, "توزیع زمان انتظار", ds, dist_len, "D1")

        if has_t:
            rc += 2
            t = df["_talk_seconds"].dropna()
            rc = self._write_stats(ws, t, "آمار مکالمه (ثانیه)", rc)
            rc += 1
            rc, ts, tdist_len = self._write_dist(
                ws, t, rc,
                bins=[0, 60, 120, 180, 300, 600, float("inf")],
                labels=[
                    "0-1m", "1-2m", "2-3m", "3-5m", "5-10m", "10m+",
                ],
            )
            self._add_bar(ws, "توزیع مدت مکالمه", ts, tdist_len, "D29")

        self.style.auto_fit(ws)

    # ── helpers ──

    def _write_stats(
        self, ws: Worksheet, series: pd.Series, title: str, rc: int
    ) -> int:
        ws.cell(row=rc, column=1, value=title)
        ws.cell(row=rc, column=2, value="مقدار")
        self.style.apply_header(ws, row=rc, max_col=2)
        rc += 1
        for k, v in {
            "میانگین": series.mean(),
            "میانه": series.median(),
            "std": series.std(),
            "min": series.min(),
            "max": series.max(),
            "p25": series.quantile(0.25),
            "p75": series.quantile(0.75),
            "p90": series.quantile(0.90),
        }.items():
            ws.cell(row=rc, column=1, value=k)
            ws.cell(row=rc, column=2, value=round(v, 1))
            rc += 1
        return rc

    def _write_dist(
        self,
        ws: Worksheet,
        series: pd.Series,
        rc: int,
        bins: list,
        labels: list,
    ) -> tuple[int, int, int]:
        dist = pd.cut(series, bins=bins, labels=labels, right=False)
        dist = dist.value_counts().sort_index().reset_index()
        dist.columns = ["بازه", "تعداد"]

        ws.cell(row=rc, column=1, value="بازه")
        ws.cell(row=rc, column=2, value="تعداد")
        self.style.apply_header(ws, row=rc, max_col=2)
        ds = rc
        rc += 1
        for _, rd in dist.iterrows():
            ws.cell(row=rc, column=1, value=rd["بازه"])
            ws.cell(row=rc, column=2, value=int(rd["تعداد"]))
            rc += 1
        return rc, ds, len(dist)

    def _add_bar(
        self, ws: Worksheet, title: str, ds: int, n: int, anchor: str
    ):
        bar = BarChart()
        bar.title = title
        bar.style = 10
        bar.width, bar.height = 25, 14
        bar.add_data(
            Reference(ws, min_col=2, min_row=ds, max_row=ds + n),
            titles_from_data=True,
        )
        bar.set_categories(
            Reference(ws, min_col=1, min_row=ds + 1, max_row=ds + n)
        )
        ws.add_chart(bar, anchor)
