"""
شیت تحلیل دلتا (فاصله زمانی مچینگ)
- دلتای پشتیبانی (Support)
- دلتای (Rate / پنل)
"""

from __future__ import annotations

from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..data_preparator import PreparedData
from .base import SheetCreator


class DeltaSheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "تحلیل_دلتا"

    # ─── ثابت‌ها ───
    BINS = [0, 1, 2, 5, 10, 15, 30, 60, float("inf")]
    BIN_LABELS = [
        "0-1m", "1-2m", "2-5m", "5-10m",
        "10-15m", "15-30m", "30-60m", "60m+",
    ]

    # ─── استایل عنوان توزیع (مرج‌شده + center) ───
    _DIST_TITLE_FONT = Font(name="B Nazanin", size=13, bold=True, color="FFFFFF")
    _DIST_TITLE_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    _DIST_TITLE_ALIGN = Alignment(horizontal="center", vertical="center")
    _DIST_TITLE_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df

        # ── بررسی وجود حداقل یک دلتا ──
        has_support = "_delta_support_minutes" in df.columns
        has_rate = "_delta_rate_minutes" in df.columns

        # ⬇️ سازگاری با نسخه قبلی
        has_legacy = "_delta_minutes" in df.columns

        if not has_support and not has_rate and not has_legacy:
            ws.cell(row=1, column=1, value="داده Delta موجود نیست")
            return

        rc = 1  # ردیف جاری

        # ═══════════════════════════════════════
        # ① دلتای پشتیبانی (Support)
        # ═══════════════════════════════════════
        if has_support:
            d_sup = df["_delta_support_minutes"].dropna()
            if not d_sup.empty:
                rc = self._write_delta_section(
                    ws, d_sup, rc,
                    title_stats="📞 آمار دلتا — پشتیبانی (دقیقه)",
                    title_dist="📞 توزیع دلتا — پشتیبانی",
                    chart_title="توزیع فاصله زمانی مچینگ پشتیبانی",
                    chart_anchor="D1",
                )
                rc += 2  # فاصله بین دو بخش

        # ═══════════════════════════════════════
        # ② دلتای (Rate / پنل)
        # ═══════════════════════════════════════
        if has_rate:
            d_rate = df["_delta_rate_minutes"].dropna()
            if not d_rate.empty:
                chart_row = rc
                rc = self._write_delta_section(
                    ws, d_rate, rc,
                    title_stats="⭐ آمار دلتا — پنل (دقیقه)",
                    title_dist="⭐ توزیع دلتا — پنل",
                    chart_title="توزیع فاصله زمانی مچینگ پنل",
                    chart_anchor=f"D{29}",
                )

        # ═══════════════════════════════════════
        # ③ سازگاری — فقط نسخه قبلی
        # ═══════════════════════════════════════
        if not has_support and not has_rate and has_legacy:
            d_old = df["_delta_minutes"].dropna()
            if not d_old.empty:
                rc = self._write_delta_section(
                    ws, d_old, rc,
                    title_stats="آمار دلتا (دقیقه)",
                    title_dist="توزیع دلتا",
                    chart_title="توزیع فاصله زمانی مچینگ",
                    chart_anchor="D1",
                )

        self.style.auto_fit(ws)

    # ──────────────────────────────────────────
    # استایل‌دهی عنوان توزیع (مرج + center)
    # ──────────────────────────────────────────
    def _style_dist_title(self, ws, row: int, max_col: int = 2) -> None:
        """
        عنوان توزیع را مرج می‌کند و استایل center + رنگی اعمال می‌کند.
        """
        ws.merge_cells(
            start_row=row, start_column=1,
            end_row=row, end_column=max_col,
        )
        cell = ws.cell(row=row, column=1)
        cell.font = self._DIST_TITLE_FONT
        cell.fill = self._DIST_TITLE_FILL
        cell.alignment = self._DIST_TITLE_ALIGN
        cell.border = self._DIST_TITLE_BORDER

        # استایل border برای سلول‌های مرج‌شده (سلول دوم هم border بگیرد)
        for col in range(2, max_col + 1):
            c = ws.cell(row=row, column=col)
            c.border = self._DIST_TITLE_BORDER

    # ──────────────────────────────────────────
    # یک بخش کامل دلتا (آمار + توزیع + نمودار)
    # ──────────────────────────────────────────
    def _write_delta_section(
        self,
        ws,
        series: pd.Series,
        start_row: int,
        title_stats: str,
        title_dist: str,
        chart_title: str,
        chart_anchor: str,
    ) -> int:
        """
        یک بخش کامل (جدول آمار + جدول توزیع + نمودار) را می‌نویسد.
        مقدار برگشتی: شماره ردیف بعد از آخرین ردیف نوشته‌شده.
        """
        rc = start_row

        # ── ① جدول آمار ──
        ws.cell(row=rc, column=1, value=title_stats)
        ws.cell(row=rc, column=2, value="مقدار")
        self.style.apply_header(ws, row=rc, max_col=2)
        rc += 1

        stats_start = rc
        for k, v in {
            "میانگین": series.mean(),
            "میانه": series.median(),
            "انحراف معیار": series.std(),
            "حداقل": series.min(),
            "حداکثر": series.max(),
            "صدک ۱۰": series.quantile(0.1),
            "صدک ۲۵": series.quantile(0.25),
            "صدک ۷۵": series.quantile(0.75),
            "صدک ۹۰": series.quantile(0.9),
            "صدک ۹۵": series.quantile(0.95),
        }.items():
            ws.cell(row=rc, column=1, value=k)
            ws.cell(row=rc, column=2, value=round(v, 2))
            rc += 1
        stats_end = rc - 1

        # ✅ border جدول آمار
        self.style.style_data(
            ws,
            start_row=stats_start,
            end_row=stats_end,
            max_col=2,
        )

        # ── ② عنوان توزیع (مرج + center + استایل) ──
        rc += 1
        ws.cell(row=rc, column=1, value=title_dist)
        self._style_dist_title(ws, row=rc, max_col=2)  # ✅ مرج + center + رنگ
        rc += 1

        # ── هدر ستون‌های توزیع ──
        ws.cell(row=rc, column=1, value="بازه")
        ws.cell(row=rc, column=2, value="تعداد")
        self.style.apply_header(ws, row=rc, max_col=2)
        ds = rc  # ردیف هدر توزیع (برای نمودار)
        rc += 1

        dist = pd.cut(
            series.abs(), bins=self.BINS,
            labels=self.BIN_LABELS, right=False,
        )
        dist = dist.value_counts().sort_index().reset_index()
        dist.columns = ["بازه", "تعداد"]

        dist_start = rc
        for _, rd in dist.iterrows():
            ws.cell(row=rc, column=1, value=rd["بازه"])
            ws.cell(row=rc, column=2, value=int(rd["تعداد"]))
            rc += 1
        dist_end = rc - 1

        # ✅ border جدول توزیع
        self.style.style_data(
            ws,
            start_row=dist_start,
            end_row=dist_end,
            max_col=2,
        )

        # ── ③ نمودار Bar ──
        bar = BarChart()
        bar.title = chart_title
        bar.style = 10
        bar.width, bar.height = 25, 14

        bar.add_data(
            Reference(ws, min_col=2, min_row=ds, max_row=ds + len(dist)),
            titles_from_data=True,
        )
        bar.set_categories(
            Reference(ws, min_col=1, min_row=ds + 1, max_row=ds + len(dist)),
        )

        # ✅ محور Y — نمایش اعداد
        bar.y_axis.delete = False
        bar.y_axis.numFmt = "#,##0"
        bar.y_axis.tickLblPos = "low"
        bar.y_axis.title = "تعداد"

        # ✅ محور X
        bar.x_axis.title = "بازه زمانی (دقیقه)"
        bar.x_axis.tickLblPos = "low"

        ws.add_chart(bar, chart_anchor)

        return rc
