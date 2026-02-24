"""
شیت تحلیل مچینگ
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList

from ..data_preparator import PreparedData
from .base import SheetCreator


class MatchingSheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "تحلیل_مچینگ"

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df
        rc = 1

        # ═══════════════════════════════════════
        # ① جدول وضعیت هر نوع مچ
        # ═══════════════════════════════════════
        ws.cell(row=rc, column=1, value="نوع")
        ws.cell(row=rc, column=2, value="وضعیت")
        ws.cell(row=rc, column=3, value="تعداد")
        ws.cell(row=rc, column=4, value="درصد")
        self.style.apply_header(ws, row=rc, max_col=4)
        rc += 1

        table1_data_start = rc

        for col, label in [
            (data.support_match_col, "پشتیبانی"),
            (data.rate_match_col, "پنل"),
        ]:
            if col and col in df.columns:
                vc = (
                    df[col]
                    .fillna("نامشخص")
                    .astype(str)
                    .value_counts()
                )
                for st, cnt in vc.items():
                    ws.cell(row=rc, column=1, value=label)
                    ws.cell(row=rc, column=2, value=st)
                    ws.cell(row=rc, column=3, value=cnt)
                    ws.cell(
                        row=rc,
                        column=4,
                        value=f"{cnt / data.total * 100:.1f}%",
                    )
                    rc += 1

        # ✅ استایل جدول اول — با end_row و max_col صریح
        table1_data_end = rc - 1
        if table1_data_end >= table1_data_start:
            self.style.style_data(
                ws,
                start_row=table1_data_start,
                end_row=table1_data_end,
                max_col=4,
            )

        # ═══════════════════════════════════════
        # ② جدول وضعیت ترکیبی
        # ═══════════════════════════════════════
        s_mask = self._match_mask(df, data.support_match_col)
        r_mask = self._match_mask(df, data.rate_match_col)

        rc += 1
        combo_header_row = rc
        ws.cell(row=rc, column=1, value="وضعیت ترکیبی")
        ws.cell(row=rc, column=2, value="تعداد")
        self.style.apply_header(ws, row=rc, max_col=2)
        rc += 1

        if s_mask.any() or r_mask.any():
            combos = [
                ("هر دو مچ", int((s_mask & r_mask).sum())),
                ("فقط پشتیبانی", int((s_mask & ~r_mask).sum())),
                ("فقط پنل", int((~s_mask & r_mask).sum())),
                ("هیچکدام", int((~s_mask & ~r_mask).sum())),
            ]
        else:
            combos = [("داده کافی نیست", 0)]

        combo_data_start = rc
        for lb, vl in combos:
            ws.cell(row=rc, column=1, value=lb)
            ws.cell(row=rc, column=2, value=vl)
            rc += 1

        # ✅ استایل جدول ترکیبی — border + font + alignment
        combo_data_end = rc - 1
        if combo_data_end >= combo_data_start:
            self.style.style_data(
                ws,
                start_row=combo_data_start,
                end_row=combo_data_end,
                max_col=2,
            )

        # ═══════════════════════════════════════
        # ③ Doughnut Chart
        # ═══════════════════════════════════════
        if len(combos) > 1:
            dnt = DoughnutChart()
            dnt.title = "توزیع ترکیبی مچینگ"
            dnt.style = 10
            dnt.width, dnt.height = 20, 14
            dnt.add_data(
                Reference(
                    ws,
                    min_col=2,
                    min_row=combo_header_row,
                    max_row=combo_data_start + len(combos) - 1,
                ),
                titles_from_data=True,
            )
            dnt.set_categories(
                Reference(
                    ws,
                    min_col=1,
                    min_row=combo_data_start,
                    max_row=combo_data_start + len(combos) - 1,
                )
            )

            # ✅ نمایش مقدار و درصد روی نمودار
            if dnt.series:
                dnt.series[0].dLbls = DataLabelList()
                dnt.series[0].dLbls.showVal = True
                dnt.series[0].dLbls.showPercent = True
                dnt.series[0].dLbls.numFmt = "#,##0"

            ws.add_chart(dnt, "F1")

        self.style.auto_fit(ws)
