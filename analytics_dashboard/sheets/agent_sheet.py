"""
شیت تحلیل اپراتور
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

from ..data_preparator import PreparedData
from .base import SheetCreator

EXCLUDED_AGENTS = {"8057", "8809", "8598"}

def _clean_agent_name(val) -> str:
    """
    ۱) NaN / خالی / "نامشخص" → "رها شده"
    ۲) اگر عددی باشد .0 آخرش حذف شود  (مثلاً 1001.0 → 1001)
    """
    if pd.isna(val) or str(val).strip() in ("", "نامشخص", "nan", "None"):
        return "رها شده"

    s = str(val).strip()

    # حذف .0 از انتهای رشته  (مثلاً "1001.0" → "1001")
    if s.endswith(".0"):
        try:
            s = str(int(float(s)))
        except (ValueError, OverflowError):
            pass

    return s


class AgentSheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "تحلیل_اپراتور"

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df

        if "_agent" not in df.columns:
            ws.cell(row=1, column=1, value="داده اپراتور موجود نیست")
            return

        # ── پاک‌سازی نام اپراتور: رفع .0 و تبدیل نامشخص → رها شده ──
        df = df.copy()
        df["_agent"] = df["_agent"].apply(_clean_agent_name)
        df = df[~df["_agent"].isin(EXCLUDED_AGENTS)]

        stats = df.groupby("_agent").agg(
            تعداد=("_agent", "size")
        ).reset_index()
        stats.rename(columns={"_agent": "اپراتور"}, inplace=True)

        if "_talk_seconds" in df.columns:
            ta = df.groupby("_agent")["_talk_seconds"].mean()
            stats["میانگین_مکالمه"] = stats["اپراتور"].map(ta).round(1)

        s_mask = self._match_mask(df, data.support_match_col)
        if s_mask.any():
            sm = df[s_mask].groupby("_agent").size()
            stats["مچ_پشتیبانی"] = (
                stats["اپراتور"].map(sm).fillna(0).astype(int)
            )

        stats.sort_values("تعداد", ascending=False, inplace=True)
        self._write_table(ws, stats)
        self.style.auto_fit(ws)

        ncols = len(stats.columns)
        n = min(len(stats), 30)

        # ── نمودار ستونی ──
        bar = BarChart()
        bar.type = "col"
        bar.title = "تعداد تماس به تفکیک اپراتور"
        bar.style = 10
        bar.width, bar.height = 30, 15

        # ✅ محور Y: مقادیر عددی مشخص باشند
        bar.y_axis.title = "تعداد تماس"
        bar.y_axis.delete = False              # محور Y حذف نشود
        bar.y_axis.numFmt = "#,##0"            # فرمت عددی (بدون اعشار)
        bar.y_axis.tickLblPos = "low"          # لیبل‌ها نمایش داده شوند

        # ✅ محور X
        bar.x_axis.title = "اپراتور"
        bar.x_axis.delete = False
        bar.x_axis.tickLblPos = "low"

        bar.add_data(
            Reference(ws, min_col=2, min_row=1, max_row=n + 1),
            titles_from_data=True,
        )
        bar.set_categories(
            Reference(ws, min_col=1, min_row=2, max_row=n + 1)
        )
        ws.add_chart(bar, f"{get_column_letter(ncols + 2)}1")
