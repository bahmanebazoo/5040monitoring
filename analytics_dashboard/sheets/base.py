"""
SheetCreator — کلاس پایه انتزاعی (OCP / LSP)
"""

from __future__ import annotations

from abc import ABC, abstractmethod
from typing import Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..data_preparator import PreparedData
from ..style_manager import StyleManager


class SheetCreator(ABC):

    GAP_ROWS = 3  # فاصله استاندارد بین بخش‌ها

    def __init__(self, style: StyleManager):
        self.style = style

    # ── اینترفیس ──

    @property
    @abstractmethod
    def sheet_name(self) -> str: ...

    @abstractmethod
    def create(self, wb: Workbook, data: PreparedData) -> None: ...

    # ── ابزارهای مشترک ──

    def _next_row(self, current_end: int) -> int:
        """ردیف شروع بخش بعدی با فاصله استاندارد."""
        return current_end + self.GAP_ROWS

    def _write_section(
        self,
        ws: Worksheet,
        title: str,
        data_df: pd.DataFrame,
        start_row: int,
    ) -> int:
        """
        یک بخش کامل می‌نویسد:
          1) عنوان merge‌شده
          2) هدر جدول
          3) داده‌ها با border
        برمی‌گرداند: اولین ردیف خالی بعد از جدول.
        """
        ncols = len(data_df.columns)

        # ۱) عنوان merge & center
        next_r = self.style.write_section_title(
            ws, row=start_row, title=title, col_count=ncols
        )

        # ۲) هدر ستون‌ها
        for ci, col_name in enumerate(data_df.columns, 1):
            ws.cell(row=next_r, column=ci, value=col_name)
        self.style.apply_header(ws, row=next_r, max_col=ncols)
        header_row = next_r
        next_r += 1

        # ۳) داده‌ها
        data_start = next_r
        for ri, row_data in enumerate(
            data_df.itertuples(index=False), data_start
        ):
            for ci, v in enumerate(row_data, 1):
                cell = ws.cell(row=ri, column=ci)
                if isinstance(v, float) and not np.isnan(v):
                    cell.value = round(v, 2)
                elif pd.isna(v):
                    cell.value = ""
                else:
                    cell.value = v

        data_end = data_start + len(data_df) - 1
        if len(data_df) == 0:
            data_end = data_start

        # ۴) استایل داده‌ها
        self.style.style_data(
            ws,
            start_row=data_start,
            end_row=data_end,
            max_col=ncols,
        )

        return data_end + 1

    def _write_table(
        self,
        ws: Worksheet,
        data_df: pd.DataFrame,
        start_row: int = 1,
    ) -> int:
        """جدول ساده بدون عنوان merge‌شده."""
        ncols = len(data_df.columns)

        for ci, cn in enumerate(data_df.columns, 1):
            ws.cell(row=start_row, column=ci, value=cn)
        self.style.apply_header(ws, row=start_row, max_col=ncols)

        data_start = start_row + 1
        for ri, row in enumerate(
            data_df.itertuples(index=False), data_start
        ):
            for ci, v in enumerate(row, 1):
                cell = ws.cell(row=ri, column=ci)
                if isinstance(v, float) and not np.isnan(v):
                    cell.value = round(v, 2)
                elif pd.isna(v):
                    cell.value = ""
                else:
                    cell.value = v

        data_end = data_start + len(data_df) - 1
        if len(data_df) == 0:
            data_end = data_start

        self.style.style_data(
            ws,
            start_row=data_start,
            end_row=data_end,
            max_col=ncols,
        )

        return data_end + 1

    def _connected_mask(self, df: pd.DataFrame) -> pd.Series:
        if "_status" not in df.columns:
            return pd.Series(False, index=df.index)
        return df["_status"].str.contains(
            "وصل|connected|answered|ANSWERED", case=False, na=False
        )

    def _match_mask(
        self, df: pd.DataFrame, col: Optional[str]
    ) -> pd.Series:
        if col is None or col not in df.columns:
            return pd.Series(False, index=df.index)
        s = df[col]
        if pd.api.types.is_bool_dtype(s):
            return s.fillna(False)
        return (
            s.fillna("")
            .astype(str)
            .str.lower()
            .str.strip()
            .isin(["matched", "true", "1"])
        )
