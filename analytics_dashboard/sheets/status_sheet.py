"""
شیت تحلیل وضعیت تماس‌ها
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice

from ..data_preparator import PreparedData
from .base import SheetCreator


# ── رنگ‌های ثابت سری‌ها ──
COLOR_CONNECTED = "4472C4"   # آبی — وصل شده
COLOR_ABANDONED = "FF0000"   # قرمز — رها شده


class StatusSheetCreator(SheetCreator):

    @property
    def sheet_name(self) -> str:
        return "تحلیل_وضعیت"

    def create(self, wb: Workbook, data: PreparedData) -> None:
        ws = wb.create_sheet(self.sheet_name)
        ws.sheet_view.rightToLeft = True

        df = data.df

        if "_status" not in df.columns:
            ws.cell(row=1, column=1, value="داده وضعیت موجود نیست")
            return

        vc = df["_status"].value_counts().reset_index()
        vc.columns = ["وضعیت", "تعداد"]
        vc["درصد"] = (vc["تعداد"] / vc["تعداد"].sum() * 100).round(1)

        # ═══════════════════════════════════════
        # ① جدول اصلی وضعیت (بالا)
        # ═══════════════════════════════════════
        self._write_table(ws, vc)
        self.style.auto_fit(ws)
        n = len(vc)

        # ═══════════════════════════════════════
        # ② Pie Chart
        # ═══════════════════════════════════════
        pie = PieChart()
        pie.title = "وضعیت تماس‌ها"
        pie.style = 10
        pie.width, pie.height = 20, 14
        pie.add_data(
            Reference(ws, min_col=2, min_row=1, max_row=n + 1),
            titles_from_data=True,
        )
        pie.set_categories(
            Reference(ws, min_col=1, min_row=2, max_row=n + 1)
        )
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showVal = True
        ws.add_chart(pie, "E1")

        # ═══════════════════════════════════════
        # ③ Stacked Bar — ساعتی
        # ═══════════════════════════════════════
        if "_hour" in df.columns:
            self._add_hourly_status(ws, df, n)

    # ──────────────────────────────────────────
    # جدول ساعتی + نمودار stacked bar
    # ──────────────────────────────────────────
    def _add_hourly_status(
        self, ws, df: pd.DataFrame, status_count: int
    ):
        cr = status_count + 4
        cross = pd.crosstab(df["_hour"], df["_status"])
        num_status_cols = len(cross.columns)
        total_cols = num_status_cols + 1  # ستون ساعت + ستون‌های وضعیت

        # ── هدر جدول ──
        ws.cell(row=cr, column=1, value="ساعت")
        for ci, cn in enumerate(cross.columns, 2):
            ws.cell(row=cr, column=ci, value=cn)
        self.style.apply_header(ws, row=cr, max_col=total_cols)

        # ── ردیف‌های داده ──
        data_start = cr + 1
        for ri, (hr, rd) in enumerate(cross.iterrows(), data_start):
            ws.cell(row=ri, column=1, value=int(hr))
            for ci, v in enumerate(rd, 2):
                ws.cell(row=ri, column=ci, value=int(v))
        data_end = data_start + len(cross) - 1

        # ✅ border جدول ساعتی
        self.style.style_data(
            ws,
            start_row=data_start,
            end_row=data_end,
            max_col=total_cols,
        )

        # ── ساخت نمودار Stacked Bar ──
        stk = BarChart()
        stk.type = "col"
        stk.grouping = "stacked"
        stk.title = "وضعیت به تفکیک ساعت"
        stk.style = 10
        stk.width, stk.height = 30, 15

        stk.add_data(
            Reference(
                ws,
                min_col=2,
                max_col=total_cols,
                min_row=cr,
                max_row=cr + len(cross),
            ),
            titles_from_data=True,
        )
        stk.set_categories(
            Reference(
                ws,
                min_col=1,
                min_row=cr + 1,
                max_row=cr + len(cross),
            )
        )

        # ✅ محور Y — نمایش اعداد
        stk.y_axis.delete = False
        stk.y_axis.numFmt = "#,##0"
        stk.y_axis.tickLblPos = "low"
        stk.y_axis.title = "تعداد"

        # ✅ محور X
        stk.x_axis.title = "ساعت"
        stk.x_axis.tickLblPos = "low"

        # ✅ رنگ‌بندی سری‌ها — "رها شده" قرمز، "وصل شده" آبی
        status_names = list(cross.columns)
        for idx, s in enumerate(stk.series):
            col_name = status_names[idx] if idx < len(status_names) else ""

            # تعیین رنگ بر اساس نام وضعیت
            color = self._get_status_color(col_name)
            s.graphicalProperties.solidFill = color

        ws.add_chart(stk, "E28")

    # ──────────────────────────────────────────
    # تعیین رنگ بر اساس نام وضعیت
    # ──────────────────────────────────────────
    @staticmethod
    def _get_status_color(status_name: str) -> str:
        """
        رنگ هر وضعیت را برمی‌گرداند.
        قابل گسترش برای وضعیت‌های جدید.
        """
        name = str(status_name).strip()

        # ── رها شده / بی‌پاسخ → قرمز ──
        abandoned_keywords = [
            "رها", "رهاشده", "رها شده",
            "بی‌پاسخ", "بی پاسخ", "بدون پاسخ",
            "abandon", "missed", "unanswered",
            "نامشخص",
        ]
        for kw in abandoned_keywords:
            if kw in name.lower():
                return COLOR_ABANDONED

        # ── وصل شده → آبی ──
        connected_keywords = [
            "وصل", "وصل شده", "وصل‌شده",
            "پاسخ", "answer", "connected",
        ]
        for kw in connected_keywords:
            if kw in name.lower():
                return COLOR_CONNECTED

        # ── سایر → خاکستری ──
        return "808080"
